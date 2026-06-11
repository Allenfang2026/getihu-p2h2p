#!/usr/bin/env python3
# 把说课稿 markdown 转成排版干净、方便老师在 Word 里直接改的 docx
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

src, out = sys.argv[1], sys.argv[2]
lines = open(src, encoding="utf-8").read().splitlines()

doc = Document()
# 正文默认字体
style = doc.styles["Normal"]
style.font.name = "PingFang SC"
style.font.size = Pt(14)
try:
    style.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "PingFang SC")
except Exception:
    pass

def strip_bold(t):
    return re.sub(r"\*\*(.+?)\*\*", r"\1", t)

for ln in lines:
    s = ln.rstrip()
    if not s:
        continue
    if s.startswith("> "):          # 引用元信息
        p = doc.add_paragraph(s[2:])
        for r in p.runs:
            r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x88,0x88,0x88)
        continue
    if s.startswith("# "):          # 标题
        h = doc.add_heading(s[2:], level=0)
        continue
    if s.startswith("---"):
        continue
    # 段落标记 **【环节 · 时长】**
    m = re.match(r"\*\*【(.+?)】\*\*", s)
    if m:
        p = doc.add_paragraph()
        run = p.add_run("【" + m.group(1) + "】")
        run.bold = True; run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0xC0,0x6B,0x3A)  # 暖橙
        continue
    # 普通正文
    p = doc.add_paragraph(strip_bold(s))
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)

doc.save(out)
print("saved:", out)
